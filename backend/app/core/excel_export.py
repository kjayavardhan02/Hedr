"""Generates the Burp History Security Header Analysis workbook.

Consumes the same normalized BurpAnalysisResult the UI reads - this module
performs no analysis of its own (see hedr-burp-history-ui-excel-export.md
section 34: "The Excel generator should not perform security analysis
itself.").
"""
from __future__ import annotations

import io
import re
from datetime import timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from app import models
from app.core.applicability import label as target_type_label
from app.core.scoring import get_severity
from app.schemas import BurpAnalysisResult, CSPPolicy, EndpointAnalysisOut, PolicyHeaderIn

_EM_DASH = "—"

_ORANGE = "FFE8630A"
_DARK = "FF1F2937"
_ZEBRA = "FFFBF6F1"
_SECTION_FONT = Font(bold=True, size=12, color=_ORANGE)
_TITLE_FILL = PatternFill(start_color=_DARK, end_color=_DARK, fill_type="solid")
_TITLE_FONT = Font(bold=True, size=16, color="FFFFFFFF")
_HEADER_FILL = PatternFill(start_color=_ORANGE, end_color=_ORANGE, fill_type="solid")
_HEADER_FONT = Font(color="FFFFFFFF", bold=True)
_ZEBRA_FILL = PatternFill(start_color=_ZEBRA, end_color=_ZEBRA, fill_type="solid")
_LABEL_FILL = PatternFill(start_color="FFF3E8DC", end_color="FFF3E8DC", fill_type="solid")
_THIN = Side(style="thin", color="FFD9D2C7")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
# Every cell in every table is centred, wrapping long values (e.g. a CSP) so
# they stay readable instead of being truncated.
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# (fill, font colour) per status label - the colour language the UI uses.
_STATUS_STYLES = {
    "Present": ("FFC6EFCE", "FF276221"),
    "Missing": ("FFFFC7CE", "FF9C0006"),
    "Invalid": ("FFFFEB9C", "FF9C5700"),
    "Not Applicable": ("FFE7E6E6", "FF595959"),
}
_SEVERITY_STYLES = {
    "Critical": ("FFFFC7CE", "FF9C0006"),
    "High": ("FFFFD9B3", "FF9A3C00"),
    "Medium": ("FFFFEB9C", "FF9C5700"),
    "Low": ("FFDDEBF7", "FF1F4E78"),
    "Info": ("FFE7E6E6", "FF595959"),
}


def _display_value(status: str, actual_value: str | None) -> str:
    """Missing and Not Applicable are never confused with a real value - both
    render as an em dash, never the word "Missing" itself, so a reviewer
    filtering this column for real CSP values never has to exclude that
    literal string (spec section 15)."""
    if status in ("missing", "not_applicable"):
        return _EM_DASH
    return actual_value or _EM_DASH


def _status_label(status: str) -> str:
    return {"present": "Present", "missing": "Missing", "invalid": "Invalid", "not_applicable": "Not Applicable"}[
        status
    ]


def _readable_dt(value) -> str:
    """e.g. "01 Oct 2026, 02:35 PM UTC" instead of an ISO-8601 string."""
    if not value:
        return ""
    if value.tzinfo is not None:
        value = value.astimezone(timezone.utc)
    return value.strftime("%d %b %Y, %I:%M %p UTC")


def _fill(color: str) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _autosize(ws: Worksheet, max_width: int = 60, min_width: int = 12) -> None:
    for column_cells in ws.columns:
        first = next((c for c in column_cells if c.value is not None), None)
        if first is None:
            continue
        longest = max((len(str(c.value)) for c in column_cells if c.value is not None), default=10)
        ws.column_dimensions[first.column_letter].width = min(max(longest + 4, min_width), max_width)


def _title_banner(ws: Worksheet, text: str, ncols: int) -> None:
    ws.append([text])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    cell = ws.cell(row=1, column=1)
    cell.fill = _TITLE_FILL
    cell.font = _TITLE_FONT
    cell.alignment = _CENTER
    ws.row_dimensions[1].height = 34


def _section(ws: Worksheet, text: str, ncols: int) -> None:
    ws.append([])
    ws.append([text])
    row = ws.max_row
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    cell = ws.cell(row=row, column=1)
    cell.font = _SECTION_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 22


def _style_table(
    ws: Worksheet, header_row: int, first_data_row: int, last_row: int, ncols: int, first_col: int = 1
) -> None:
    """Orange header, thin borders, zebra rows, everything centred, and
    status/severity cells colour-coded."""
    last_col = first_col + ncols - 1
    for col in range(first_col, last_col + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER
        cell.border = _BORDER
    ws.row_dimensions[header_row].height = 24
    for row in range(first_data_row, last_row + 1):
        zebra = (row - first_data_row) % 2 == 1
        for col in range(first_col, last_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.alignment = _CENTER
            cell.border = _BORDER
            if zebra:
                cell.fill = _ZEBRA_FILL
            style = _STATUS_STYLES.get(str(cell.value)) or _SEVERITY_STYLES.get(str(cell.value))
            if style:
                cell.fill = _fill(style[0])
                cell.font = Font(bold=True, color=style[1])


def _finish_data_sheet(ws: Worksheet, ncols: int, nrows: int) -> None:
    """For a sheet whose row 1 is the header and the rest is data."""
    _style_table(ws, 1, 2, nrows + 1, ncols)
    ws.freeze_panes = "A2"
    if nrows:
        ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{nrows + 1}"
    ws.sheet_view.showGridLines = False


def _write_table(ws: Worksheet, headers: list[str], rows: list[list[object]]) -> None:
    """Append a headed table at the current position of a banner-style sheet."""
    ws.append(headers)
    header_row = ws.max_row
    for r in rows:
        ws.append(r)
    _style_table(ws, header_row, header_row + 1, ws.max_row, len(headers))


def _key_values(ws: Worksheet, pairs: list[tuple[str, object]], merge_to: int | None = None) -> None:
    """Label/value rows. `merge_to` stretches each value across to that
    column, so a long value on a wide sheet doesn't force column B wide."""
    for label, value in pairs:
        ws.append([label, value])
        row = ws.max_row
        a, b = ws.cell(row=row, column=1), ws.cell(row=row, column=2)
        a.fill = _LABEL_FILL
        a.font = Font(bold=True)
        for c in (a, b):
            c.alignment = _CENTER
            c.border = _BORDER
        if merge_to:
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=merge_to)
            for col in range(3, merge_to + 1):
                ws.cell(row=row, column=col).border = _BORDER


# ---------------------------------------------------------------------------
# Summary
# ---------------------------------------------------------------------------

_LEGEND = [
    ("Present", "The header was sent and meets the policy."),
    ("Missing", "The policy requires this header but the response did not send it. Counts as a failed check."),
    ("Invalid", "The header was sent but its value breaks the policy. Counts as a failed check."),
    (
        "Not Applicable",
        "The header does not apply to this response (e.g. HSTS on plain HTTP, CSP on a non-HTML response). "
        "Left out of the pass/fail counts and the score.",
    ),
    (_EM_DASH + " (dash)", "No value to show: the header is Missing or Not Applicable, so there is nothing to display."),
    (
        "Policy Score",
        "Per response: the percentage of applicable policy checks that passed (Not Applicable checks are "
        "excluded). The overall score is the average.",
    ),
    ("Coverage", "Per header: Present / (Present + Missing + Invalid), as a percentage. Not Applicable is excluded."),
    ("Critical / High / Medium / Low / Info", "How serious a Missing or Invalid finding is for that header."),
]


def _legend_block(ws: Worksheet, start_row: int) -> None:
    """The "What the labels mean" table, down the right-hand side of the
    Summary (columns D-E), level with the left-hand blocks."""
    ws.merge_cells(start_row=start_row, start_column=4, end_row=start_row, end_column=5)
    title = ws.cell(row=start_row, column=4, value="What the labels mean")
    title.font = _SECTION_FONT
    title.alignment = Alignment(horizontal="left", vertical="center")

    header_row = start_row + 1
    ws.cell(row=header_row, column=4, value="Label")
    ws.cell(row=header_row, column=5, value="Meaning")
    for i, (label, meaning) in enumerate(_LEGEND, start=1):
        ws.cell(row=header_row + i, column=4, value=label)
        ws.cell(row=header_row + i, column=5, value=meaning)
    _style_table(ws, header_row, header_row + 1, header_row + len(_LEGEND), 2, first_col=4)
    for row in range(header_row + 1, header_row + len(_LEGEND) + 1):
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 0, 32)


def _summary_sheet(wb: Workbook, record: models.BurpImport, analysis: BurpAnalysisResult) -> None:
    ws = wb.create_sheet("Summary", 0)
    _title_banner(ws, "Burp History Security Header Analysis", 5)

    endpoints = analysis.endpoints
    total_responses = len(endpoints)
    failed_responses = sum(1 for e in endpoints if e.has_findings)
    passed_responses = total_responses - failed_responses

    present = sum(r.present for r in analysis.header_coverage)
    failed_checks = sum(r.missing + r.invalid for r in analysis.header_coverage)
    total_checks = present + failed_checks
    not_applicable_checks = sum(r.not_applicable for r in analysis.header_coverage)

    _section(ws, "Report Details", 2)
    details_row = ws.max_row
    _key_values(
        ws,
        [
            ("Report Name", record.name),
            ("Source File", record.source_filename),
            ("Policy", record.policy_name or ""),
            ("Policy Version", record.policy_version or ""),
            ("Analysis Type", "Burp History"),
            ("Target Type", target_type_label(analysis.target_type)),
            ("Imported At", _readable_dt(record.imported_at)),
            ("Analyzed At", _readable_dt(record.analyzed_at)),
        ],
    )

    _section(ws, "Results", 2)
    _key_values(
        ws,
        [
            ("Overall Policy Score", f"{analysis.summary.overall_score}%"),
            ("Responses Passed", f"{passed_responses} of {total_responses}"),
            ("Responses Failed", f"{failed_responses} of {total_responses}"),
            ("Header Checks Passed", f"{present} of {total_checks}"),
            ("Header Checks Failed", f"{failed_checks} of {total_checks}"),
            ("Header Checks Not Applicable", not_applicable_checks),
        ],
    )
    ws.append(
        [
            "A response passes when every applicable header on it meets the policy; "
            "it fails if any header is Missing or Invalid.",
        ]
    )
    note_row = ws.max_row
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=2)
    ws.cell(row=note_row, column=1).font = Font(italic=True, color="FF6B7280")
    ws.cell(row=note_row, column=1).alignment = _CENTER
    ws.row_dimensions[note_row].height = 32

    _section(ws, "Import Metrics", 2)
    _write_table(
        ws,
        ["Metric", "Value"],
        [
            ["Total History Entries", record.entries_found],
            ["Responses Analyzed", analysis.summary.responses_analyzed],
            ["Responses Skipped", record.responses_skipped],
            ["Parse Failures", record.parse_failures],
            ["Unique Domains", analysis.summary.unique_hosts],
            ["Unique Paths", analysis.summary.unique_paths],
            ["Responses With Findings", analysis.summary.responses_with_findings],
        ],
    )

    _legend_block(ws, start_row=details_row)

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 16  # spacer between the two halves
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 80
    ws.sheet_view.showGridLines = False


# ---------------------------------------------------------------------------
# Policy
# ---------------------------------------------------------------------------


def _join(items: list[str] | None) -> str:
    return ", ".join(items) if items else _EM_DASH


def _csp_restrictions(rule) -> str:
    flags = [
        ("No wildcards (*)", rule.disallow_wildcards),
        ("No external sources", rule.disallow_external),
        ("No http: sources", rule.disallow_http),
        ("No data: sources", rule.disallow_data),
        ("No blob: sources", rule.disallow_blob),
    ]
    return ", ".join(label for label, on in flags if on) or _EM_DASH


def _fallback_headers(analysis: BurpAnalysisResult) -> list[PolicyHeaderIn]:
    """Policy rules recovered from the stored analysis, for when the saved
    policy has since been deleted or edited: the header names and expected
    values every endpoint was judged against."""
    seen: dict[str, str] = {}
    for endpoint in analysis.endpoints:
        for hr in endpoint.header_results:
            if hr.header != "Content-Security-Policy" and hr.header not in seen:
                seen[hr.header] = hr.expected_value or ""
    return [PolicyHeaderIn(header_name=n, expected_value=v, required=True) for n, v in seen.items()]


def _policy_sheet(
    wb: Workbook,
    record: models.BurpImport,
    analysis: BurpAnalysisResult,
    policy: models.Policy | None,
) -> None:
    ws = wb.create_sheet("Policy", 1)
    _title_banner(ws, "Policy Used For This Analysis", 5)

    # Only trust the saved policy if it is still the version that was analyzed.
    exact = policy is not None and f"v{policy.version}" == record.policy_version
    if exact:
        headers = [PolicyHeaderIn(**h) for h in policy.headers]
        csp = CSPPolicy(**policy.csp_policy) if policy.csp_policy else None
        description = policy.description or ""
        note = None
    else:
        headers = _fallback_headers(analysis)
        csp = None
        description = ""
        if policy is None:
            note = "The saved policy no longer exists, so the rules below are rebuilt from the analysis itself (CSP rules are not available)."
        else:
            note = (
                f"The policy has been edited since this analysis (now v{policy.version}); "
                "the rules below are rebuilt from the analysis itself (CSP rules are not available)."
            )

    _section(ws, "Policy Details", 5)
    _key_values(
        ws,
        [
            ("Policy Name", record.policy_name or ""),
            ("Version", record.policy_version or ""),
            ("Description", description or _EM_DASH),
            ("Header Rules", len(headers)),
            ("Content-Security-Policy", "Evaluated" if csp else "Not evaluated"),
        ],
        merge_to=5,
    )
    if note:
        ws.append([note])
        row = ws.max_row
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        ws.cell(row=row, column=1).font = Font(italic=True, color="FF9C0006")
        ws.cell(row=row, column=1).alignment = _CENTER
        ws.row_dimensions[row].height = 32

    _section(ws, "Header Rules", 5)
    rows = [
        [h.header_name, h.expected_value or _EM_DASH, "Yes" if h.required else "No", get_severity(h.header_name).capitalize()]
        for h in headers
    ]
    if rows:
        _write_table(ws, ["Header", "Expected Value", "Required", "Severity"], rows)
    else:
        ws.append(["No header rules."])
        ws.cell(row=ws.max_row, column=1).alignment = _CENTER

    if csp:
        _section(ws, "Content-Security-Policy Rules", 5)
        _key_values(
            ws,
            [
                ("CSP Header Required", "Yes" if csp.required else "No"),
                ("Required Directives", _join(csp.required_directives)),
            ],
            merge_to=5,
        )
        if csp.directive_rules:
            ws.append([])
            _write_table(
                ws,
                ["Directive", "Must Contain", "Must Not Contain", "Allowed Sources", "Restrictions"],
                [
                    [
                        r.directive,
                        _join(r.must_contain),
                        _join(r.must_not_contain),
                        _join(r.allowed_sources),
                        _csp_restrictions(r),
                    ]
                    for r in csp.directive_rules
                ],
            )

    ws.column_dimensions["A"].width = 32
    for col in "BCDE":
        ws.column_dimensions[col].width = 34
    ws.sheet_view.showGridLines = False


# ---------------------------------------------------------------------------
# Data sheets
# ---------------------------------------------------------------------------


def _endpoint_analysis_sheet(wb: Workbook, record: models.BurpImport, analysis: BurpAnalysisResult) -> None:
    ws = wb.create_sheet("Endpoint Analysis")
    header_names = [row.header for row in analysis.header_coverage]

    columns = ["Domain", "Path", "Method", "Status Code", "Content-Type"]
    for name in header_names:
        columns += [name, f"{name} Status"]
    columns += ["Policy Score", "Findings", "Not Applicable (Reason)"]
    ws.append(columns)

    for endpoint in analysis.endpoints:
        by_header = {hr.header: hr for hr in endpoint.header_results}
        row: list[object] = [
            endpoint.domain,
            endpoint.path,
            endpoint.method,
            endpoint.status_code,
            endpoint.content_type or "",
        ]
        failing: list[str] = []
        not_applicable: list[str] = []
        for name in header_names:
            hr = by_header.get(name)
            if hr is None:
                row += [_EM_DASH, _EM_DASH]
                continue
            row += [_display_value(hr.status, hr.actual_value), _status_label(hr.status)]
            if hr.status in ("missing", "invalid"):
                failing.append(f"{name} {hr.status}")
            elif hr.status == "not_applicable":
                not_applicable.append(f"{name}: {hr.applicability_reason or 'Not applicable.'}")
        row += [
            f"{endpoint.policy_score}%",
            "; ".join(failing) if failing else "None",
            "\n".join(not_applicable) if not_applicable else "None",
        ]
        ws.append(row)

    _finish_data_sheet(ws, len(columns), len(analysis.endpoints))
    _autosize(ws)
    # Long CSP-style values need to stay readable rather than truncated.
    for name in header_names:
        col_idx = columns.index(name) + 1
        ws.column_dimensions[get_column_letter(col_idx)].width = 40


def _header_coverage_sheet(wb: Workbook, analysis: BurpAnalysisResult) -> None:
    ws = wb.create_sheet("Header Coverage")
    ws.append(["Header", "Present", "Missing", "Invalid", "Not Applicable", "Coverage"])
    for row in analysis.header_coverage:
        ws.append([row.header, row.present, row.missing, row.invalid, row.not_applicable, f"{row.coverage}%"])
    _finish_data_sheet(ws, 6, len(analysis.header_coverage))
    _autosize(ws)


def _findings_sheet(wb: Workbook, record: models.BurpImport, analysis: BurpAnalysisResult) -> None:
    ws = wb.create_sheet("Findings")
    columns = [
        "Domain",
        "Path",
        "Method",
        "Status Code",
        "Header",
        "Finding",
        "Severity",
        "Expected Value",
        "Actual Value",
        "Policy",
        "Policy Version",
    ]
    ws.append(columns)

    def _endpoint_findings(endpoint: EndpointAnalysisOut):
        for hr in endpoint.header_results:
            if hr.status not in ("missing", "invalid"):
                continue
            yield hr

    nrows = 0
    for endpoint in analysis.endpoints:
        for hr in _endpoint_findings(endpoint):
            ws.append(
                [
                    endpoint.domain,
                    endpoint.path,
                    endpoint.method,
                    endpoint.status_code,
                    hr.header,
                    _status_label(hr.status),
                    (hr.severity or "").capitalize(),
                    hr.expected_value or _EM_DASH,
                    _display_value(hr.status, hr.actual_value),
                    record.policy_name or "",
                    record.policy_version or "",
                ]
            )
            nrows += 1

    _finish_data_sheet(ws, len(columns), nrows)
    _autosize(ws)


def _import_issues_sheet(wb: Workbook, analysis: BurpAnalysisResult) -> None:
    if not analysis.import_issues:
        return
    ws = wb.create_sheet("Import Issues")
    ws.append(["Entry", "URL", "Host", "Path", "Status", "Reason"])
    for issue in analysis.import_issues:
        ws.append([issue.index, issue.url or "", issue.host or "", issue.path or "", issue.status.capitalize(), issue.reason or ""])
    _finish_data_sheet(ws, 6, len(analysis.import_issues))
    _autosize(ws)


def build_workbook(record: models.BurpImport, policy: models.Policy | None = None) -> bytes:
    """`policy` is the saved policy the import was analyzed against, if it
    still exists - the Policy sheet lists its full rule set (see _policy_sheet
    for what happens when it is missing or has been edited since)."""
    if record.analysis is None:
        raise ValueError("This import hasn't been analyzed yet.")
    analysis = BurpAnalysisResult.model_validate(record.analysis)

    wb = Workbook()
    wb.remove(wb.active)  # the default blank sheet - each sheet below is created explicitly, in order
    _summary_sheet(wb, record, analysis)
    _policy_sheet(wb, record, analysis, policy)
    _endpoint_analysis_sheet(wb, record, analysis)
    _header_coverage_sheet(wb, analysis)
    _findings_sheet(wb, record, analysis)
    _import_issues_sheet(wb, analysis)

    for ws in wb.worksheets:
        # Print/PDF friendly: landscape, one page wide.
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_filename(record: models.BurpImport) -> str:
    # Strip characters that are unsafe in a filename across platforms/browsers.
    safe_name = re.sub(r'[\\/*?:"<>|]', "", record.name).strip() or "Burp History"
    return f"{safe_name} - Security Header Analysis.xlsx"
