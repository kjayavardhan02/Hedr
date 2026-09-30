import io

import openpyxl
import pytest

from app.core.excel_export import build_workbook, export_filename
from app.core.policy_resolution import ResolvedPolicy
from app.schemas import CSPPolicy, PolicyHeaderIn
from app.core.burp_aggregation import build_analysis
from app.core.burp_import import BurpEntry, BurpEntryRequest, BurpEntryResponse, EntryStatus


class FakeRecord:
    """A minimal stand-in for models.BurpImport - build_workbook only reads
    plain attributes, never touches the DB, so a real ORM object isn't needed."""

    def __init__(self, analysis_model, **overrides):
        self.name = overrides.get("name", "Test Assessment")
        self.source_filename = overrides.get("source_filename", "history.xml")
        self.policy_name = overrides.get("policy_name", "Strict")
        self.policy_version = overrides.get("policy_version", "v1")
        self.entries_found = overrides.get("entries_found", 3)
        self.responses_skipped = overrides.get("responses_skipped", 0)
        self.parse_failures = overrides.get("parse_failures", 0)
        self.imported_at = overrides.get("imported_at")
        self.analyzed_at = overrides.get("analyzed_at")
        self.analysis = analysis_model.model_dump(mode="json") if analysis_model is not None else None


HSTS_HEADER = PolicyHeaderIn(header_name="Strict-Transport-Security", expected_value="max-age=31536000", required=True)


def _entry(index, *, url, host, path="/", headers=None, content_type="text/html"):
    request = BurpEntryRequest(method="GET", url=url, host=host, port=443, path=path)
    response = BurpEntryResponse(status_code=200, headers=headers or {}, content_type=content_type)
    return BurpEntry(index=index, request=request, response=response, status=EntryStatus.PARSED)


def _sample_analysis():
    resolved = ResolvedPolicy(
        policy=None,  # type: ignore[arg-type]
        policy_headers=[HSTS_HEADER],
        policy_name="Strict",
        policy_version="v1",
        csp_policy=CSPPolicy(required=True),
    )
    e1 = _entry(1, url="https://a.com/", host="a.com", headers={"strict-transport-security": "max-age=31536000"})
    e2 = _entry(2, url="https://b.com/x", host="b.com", path="/x", headers={})
    return build_analysis([e1, e2], resolved)


def _load(xlsx_bytes: bytes):
    return openpyxl.load_workbook(io.BytesIO(xlsx_bytes))


def test_workbook_has_expected_sheets_in_order():
    record = FakeRecord(_sample_analysis())
    wb = _load(build_workbook(record))
    assert wb.sheetnames == ["Endpoint Analysis", "Summary", "Header Coverage", "Findings"]


def test_endpoint_analysis_sheet_content():
    record = FakeRecord(_sample_analysis())
    wb = _load(build_workbook(record))
    ws = wb["Endpoint Analysis"]
    header_row = [c.value for c in ws[1]]
    assert "Domain" in header_row
    assert "Strict-Transport-Security" in header_row
    assert "Strict-Transport-Security Status" in header_row
    assert "Content-Security-Policy" in header_row

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    by_domain = {r[0]: r for r in rows}
    hsts_col = header_row.index("Strict-Transport-Security")
    hsts_status_col = header_row.index("Strict-Transport-Security Status")

    # a.com has HSTS present with its real value.
    assert by_domain["a.com"][hsts_col] == "max-age=31536000"
    assert by_domain["a.com"][hsts_status_col] == "Present"
    # b.com never sent it - the value cell is an em dash, never blank or "Missing".
    assert by_domain["b.com"][hsts_col] == "—"
    assert by_domain["b.com"][hsts_status_col] == "Missing"


def test_missing_and_not_applicable_never_leak_the_status_word_into_the_value_cell():
    record = FakeRecord(_sample_analysis())
    wb = _load(build_workbook(record))
    ws = wb["Endpoint Analysis"]
    header_row = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        for col_name in header_row:
            if not col_name or col_name.endswith(" Status") or col_name in (
                "Domain", "Path", "Method", "Status Code", "Content-Type", "Policy Score", "Findings"
            ):
                continue
            idx = header_row.index(col_name)
            value = row[idx]
            assert value != "Missing"
            assert value != "Not Applicable"


def test_summary_sheet_metrics():
    record = FakeRecord(_sample_analysis(), entries_found=5, responses_skipped=1, parse_failures=2)
    wb = _load(build_workbook(record))
    ws = wb["Summary"]
    values = {row[0]: row[1] for row in ws.iter_rows(values_only=True) if row[0]}
    assert values["Report Name"] == "Test Assessment"
    assert values["Policy"] == "Strict"
    assert values["Total History Entries"] == 5
    assert values["Responses Skipped"] == 1
    assert values["Parse Failures"] == 2
    assert values["Responses Analyzed"] == 2
    assert values["Unique Domains"] == 2


def test_header_coverage_sheet():
    record = FakeRecord(_sample_analysis())
    wb = _load(build_workbook(record))
    ws = wb["Header Coverage"]
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert "Strict-Transport-Security" in rows
    present, missing, invalid, na, coverage = rows["Strict-Transport-Security"][1:]
    assert present == 1
    assert missing == 1
    assert coverage == "50.0%"


def test_findings_sheet_has_one_row_per_failing_header():
    record = FakeRecord(_sample_analysis())
    wb = _load(build_workbook(record))
    ws = wb["Findings"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    # b.com is missing HSTS and CSP (CSP applies to its HTML response) - two finding rows.
    domains = [r[0] for r in rows]
    assert domains.count("b.com") == 2
    headers = {r[4] for r in rows if r[0] == "b.com"}
    assert headers == {"Strict-Transport-Security", "Content-Security-Policy"}


def test_import_issues_sheet_only_present_when_there_are_issues():
    record = FakeRecord(_sample_analysis())
    wb = _load(build_workbook(record))
    assert "Import Issues" not in wb.sheetnames


def test_not_analyzed_raises():
    record = FakeRecord(None)
    with pytest.raises(ValueError):
        build_workbook(record)


def test_export_filename_sanitizes_unsafe_characters():
    record = FakeRecord(_sample_analysis(), name='Weird/Name:With*Bad?Chars"<>|')
    filename = export_filename(record)
    assert filename == "WeirdNameWithBadChars - Security Header Analysis.xlsx"


def test_export_filename_falls_back_when_name_is_empty():
    record = FakeRecord(_sample_analysis(), name="///")
    filename = export_filename(record)
    assert filename == "Burp History - Security Header Analysis.xlsx"
